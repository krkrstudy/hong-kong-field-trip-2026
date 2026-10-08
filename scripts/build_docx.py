"""Generate the editable handbook from the shared data.json. Requires python-docx, Pillow."""
import json, math, re
from pathlib import Path
from urllib.parse import quote
from PIL import Image,ImageOps
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R=Path(__file__).resolve().parents[1]; D=json.loads((R/'data.json').read_text());DOC=Document();S=DOC.sections[0]
S.page_width=Cm(21);S.page_height=Cm(29.7);S.top_margin=Cm(1.9);S.bottom_margin=Cm(1.7);S.left_margin=Cm(2);S.right_margin=Cm(2);S.header_distance=Cm(.8);S.footer_distance=Cm(.8);S.different_first_page_header_footer=True
INK='17394D';GREEN='376B63';MUTED='596973';LINE='D9DEDC';WHITE='FFFFFF'
styles=DOC.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=styles[name];st.font.name='Arial';st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'STSong' if name in ['Normal','Caption'] else 'Arial Unicode MS');st.font.color.rgb=RGBColor.from_string(INK);st.paragraph_format.line_spacing=1.5
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(30);styles['Title'].font.bold=True
styles['Subtitle'].font.italic=False
styles['Caption'].font.bold=False;styles['Caption'].font.italic=False
styles['Heading 1'].font.size=Pt(21);styles['Heading 1'].paragraph_format.space_before=Pt(0);styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Heading 2'].font.size=Pt(14);styles['Heading 2'].paragraph_format.space_before=Pt(12);styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Heading 3'].font.size=Pt(11);styles['Heading 3'].font.color.rgb=RGBColor.from_string(GREEN)
styles['Caption'].font.size=Pt(8);styles['Caption'].font.color.rgb=RGBColor.from_string(MUTED);styles['Caption'].paragraph_format.space_after=Pt(7)
for name in ['Heading 1','Heading 2','Heading 3']:styles[name].paragraph_format.keep_with_next=True
for name in ['Normal','Caption']:styles[name].paragraph_format.widow_control=True
for st in styles:
 for fonts in st.element.xpath('.//w:rFonts'):
  for key in list(fonts.attrib):
   if 'Theme' in key:del fonts.attrib[key]
head=S.header.paragraphs[0];head.text='香江翱翔  /  HONG KONG FIELD NOTES                              2026.10.11—10.17';head.paragraph_format.line_spacing=1;head.runs[0].font.size=Pt(8);head.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
foot=S.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=foot.add_run('V4 · 2026-10-07     /     ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
DOC.core_properties.title='香港访学精美路书';DOC.core_properties.subject=D['title'];DOC.core_properties.author='香港访港研学团';DOC.core_properties.keywords='香港 访学 香江翱翔 V4'
settings=DOC.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
TMP=R/'.doc-build';TMP.mkdir(exist_ok=True)
page=1;toc=[];bk=0;pending_break=False

def para(text='',style=None,size=None,color=None,bold=False):
 p=DOC.add_paragraph(text,style);p.paragraph_format.line_spacing=1.5
 if size or color or bold:
  for r in p.runs:
   if size:r.font.size=Pt(size)
   if color:r.font.color.rgb=RGBColor.from_string(color)
   if bold:r.bold=True
 return p

def title(text,eyebrow=None,toc_entry=True):
 global bk,pending_break
 if eyebrow:
  ep=para(eyebrow,size=8.5,color=GREEN,bold=True)
  if pending_break:ep.paragraph_format.page_break_before=True;pending_break=False
 p=DOC.add_paragraph(text,'Heading 1')
 if toc_entry:
  bk+=1;name='chapter'+str(bk);start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bk));start.set(qn('w:name'),name);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bk));p._p.insert(0,start);p._p.append(end);entry_text='地点深度介绍' if text=='香港国际机场' else text
  toc.append((entry_text,page,name))
  tc=OxmlElement('w:fldSimple');tc.set(qn('w:instr'),' TC "'+entry_text+'" \\f C \\l 1 ');rp=OxmlElement('w:r');props=OxmlElement('w:rPr');props.append(OxmlElement('w:vanish'));rp.append(props);tc.append(rp);p._p.append(tc)
 return p

def newpage():
 global page,pending_break
 pending_break=True;page+=1

def link(p,text,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),GREEN);pr.append(c);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def picture(path,width=17,height=None,caption=None):
 path=Path(path)
 if height:
  im=Image.open(path).convert('RGB');im=ImageOps.fit(im,(1700,int(1700*height/width)),centering=(.5,.43));path=TMP/(path.stem+f'-{height}.jpg');im.save(path,quality=94)
 p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(path),width=Cm(width));p.paragraph_format.keep_with_next=bool(caption)
 if caption:para(caption,'Caption')

def table(headers,rows,widths):
 t=DOC.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trpr.append(rep)
 for row in rows:
  cs=t.add_row().cells
  for c,x in zip(cs,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),INK if ri==0 else ('F0F4F2' if ri%2 else WHITE));pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for x,v in [('top','90'),('bottom','90'),('left','100'),('right','100')]:z=OxmlElement('w:'+x);z.set(qn('w:w'),v);z.set(qn('w:type'),'dxa');margins.append(z)
   pr.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.5
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string(WHITE if ri==0 else INK);r.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for x in ['top','left','bottom','right','insideH','insideV']:
  v=OxmlElement('w:'+x);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),LINE);borders.append(v)
 t._tbl.tblPr.append(borders)
 return t

def daytimes(p):return '；'.join('10月'+str(day['day'])+'日 '+e['time']+('—'+e['end'] if e['end'] else '') for day in D['days'] for e in day['events'] if e['place']==p['id'])
def refline(ids):
 p=para('资料核验  ','Caption')
 for i,s in enumerate(ids):
  if s=='schedule':continue
  source=next(x for x in D['sources'] if x['id']==s);link(p,source['title'],source['url']);p.add_run('  ')

# 1 Cover
para('PEKING UNIVERSITY  /  FIELD TRIP 2026',size=9,color=GREEN,bold=True)
para('香江翱翔','Title');para('香港访学与交流路书','Subtitle',size=18)
para('第五届北京大学“香江翱翔”访港研学团',size=13)
para('2026年10月11日至17日  ·  中国香港',size=11,color=MUTED)
picture(R/'assets/images/hero.jpg',height=10.8)
para('从公共治理到前沿科研\n从大学课堂到金融现场',size=17,color=INK)
para(D['organizers'],size=10)
para('依据香港校友会2026年10月7日V4正式行程编制\n信息核查日期 2026年10月8日',size=9,color=MUTED)
para('照片为香港城市实景，拍摄时间早于本次访问；版权与署名见文末。','Caption')
# 2 Contents
newpage();title('阅读导航','CONTENTS',False);toc_holder=DOC.add_paragraph()
para('使用说明','Heading 2');para('时间安排以V4为唯一依据。已知冲突和缺失信息保留“待确认”标记；领队后续通知优先于本手册。所有时间均为香港当地时间，与北京无时差。')
para('地图显示经核查的楼宇或园区位置，不代表预约入口；编号用于对应地点，不是道路路线。车程未核实的路段不提供推测分钟数。')
para('网页路书与本手册共享同一份行程数据。手机端可切换日期、阅读地点详情并打开导航；纸质版建议在行前下载，作为网络不稳定时的备用。')
link(para('数字路书  '),'krkrstudy.github.io/hong-kong-field-trip-2026/','https://krkrstudy.github.io/hong-kong-field-trip-2026/')
# 3 framing
newpage();title('在香港理解开放与创新','01 / BEFORE THE JOURNEY')
para('2026年是“十五五”开局之年。2025年12月16日，习近平主席在听取香港特别行政区行政长官李家超述职报告时，要求特区政府主动对接国家“十五五”规划，推动高质量发展，深化大湾区建设中的参与。这为青年从国家发展与香港实践的联系中开展学习提供了重要背景。')
para('香港2026年施政报告进一步讨论科研成果转化、产业发展和创新合作。此次行程将公共治理、人才服务、数字科技、生命健康和金融实践放在同一条学习路径上。参访的意义，在于把宏观发展议题还原为具体机构、工作流程和人的选择。')
para('本团由北京大学香港校友会和香港北大助学基金会发起支持，成员为来自北大多个院系的15名学生。期待同学们既看见香港连接内地与世界的条件，也认真倾听机构与个人面对的实际问题，以扎实观察形成有证据、有分寸的认识。')
para('本次学习的三条线索','Heading 2')
for text in ['制度与服务：从立法会、人才办和投资署观察公共制度如何进入城市运作。','研究与转化：从数码港、AI与生命健康研究中心观察知识如何走向应用。','专业与责任：在大学、金融机构与校友交流中，思考专业成长与社会需要的联系。']:para(text)
refline(['policy','policy2026'])
picture(R/D['overviewMap']['image'],height=6.3,caption='区域概览：1 机场；2 港岛参访区域（以上环酒店为参考）；3 科学园。底图 © OpenStreetMap contributors。')
#4 overview
newpage();title('七日行程总览','02 / THE WEEK AT A GLANCE')
rows=[]
for day in D['days']:
 events='\n'.join(e['time']+('—'+e['end'] if e['end'] else '')+' '+e['title'] for e in day['events'])
 rows.append([f"10月{day['day']}日\n{day['weekday']}",day['theme'],events])
table(['日期','每日主题','V4正式安排'],rows,[2.2,3.3,11.5])
para('港大两项活动时间重叠；爱诗地址、晚宴与校友座谈地点、中诚信房号等详见各日提醒及待确认清单。','Caption')
#5-11 daily
for idx,day in enumerate(D['days'],1):
 newpage();title(f"10月{day['day']}日  {day['theme']}",f"DAY {idx:02d} / {day['weekday']} / {day['region']}")
 rows=[]
 for e in day['events']:
  info=e['title'];
  if e['note']:info+='\n'+e['note']
  rows.append([e['time']+('—'+e['end'] if e['end'] else ''),info,e['status']])
 table(['时间','活动与地点','状态'],rows,[3.2,11,2.8])
 para('集合与交通','Heading 2');para(day['transport'],size=9.5)
 # image kept small enough for dense park day
 picture(R/day['map']['image'],height=4.8 if day['day']==15 else 5.6)
 legend='  ·  '.join(str(a['number'])+' '+a['name'] for a in day['map']['legend'])
 para(legend,'Caption');para(day['map'].get('note','编号对应地点；不连线、不推测道路路线。底图 © OpenStreetMap contributors。'),'Caption')
 para('当日任务','Heading 3');para(day['task'],size=9.5)
 para('特别留意  '+day['notice'],size=9.5,color=GREEN)
#12..25 profiles
for i,p in enumerate(D['places'],1):
 newpage();title(p['name'],f"PLACE {i:02d} / {p['category']}",toc_entry=i==1)
 para(p['en'],size=9.5,color=MUTED)
 para(daytimes(p),size=9,color=GREEN,bold=True)
 if p['image']:
  picture(R/p['image'],height=7.0,caption=p['photo']['caption'])
 else:
  para('场地照片待确认','Heading 2');para('V4没有提供到访地址，尚未找到能够确认本次接待场所的公开实景图片。为避免误导，此处不以其他地点的照片代替。')
 para('地址与入口','Heading 2');para(p['address'],size=10)
 para(p['intro'])
 para(p['topic'],'Heading 2')
 for q in p['questions']:para('• '+q,size=10)
 para('参访注意事项','Heading 2');para(p['entry'],size=10)
 if p['coordinates']:
  url='https://www.google.com/maps/search/?api=1&query='+quote(p['address']+' '+p['name']) if p['id']=='dymon' else 'https://www.google.com/maps/search/?api=1&query='+','.join(map(str,p['coordinates']))
  link(para('',style='Caption'),'打开位置导航',url)
 refline(p['sources'])
# practical
newpage();title('出行准备与参访规范','PRACTICAL NOTES')
for h,t in [
('交通与缓冲','V4已安排抵离港接送及10月15日科学园中巴。其他日期的交通方式与发车时间待领队确认。建议预约活动额外预留15—20分钟作为点名、安检和找房间的组织缓冲；这是建议，不是主办方已经确定的集合时刻。'),
('公共交通备用','港铁公布机场快线香港站至机场列车行程约24分钟，不包含酒店接驳、候车、入境及行李时间。科技园官方列有大学站转乘272K的到园方式，班次与步行衔接请临行通过港铁、香港出行易及科技园官网核实。遇团队车辆变动，先联系领队，不自行分散出发。'),
('证件与个人准备','出发前检查本人证件及适用的入境许可、机票和保险安排。随身携带常用药、雨具、饮水用品和适合步行的鞋。准备可在香港使用的通讯与支付方式，保管好证件和紧急联络信息。'),
('天气与道路安全','每日查看香港天文台天气警告；恶劣天气下遵从领队、接待方和当地部门指引。过街使用行人设施并观察来车；不要为拍照停留在车道或阻塞出入口。'),
('交流与记录','行前准备一至两个问题；现场发言简明，认真倾听。拍摄、录音和发布材料须先获许可。实验室、企业与政府机构的内部信息不得自行外传，不触碰设备或展示屏。'),
('联系与应急','私人联系方式及成员名单不放入公开路书，由领队另行发放。如遇突发情况先联系领队；紧急危险可拨打香港999求助，说明准确位置和现场情况。')]:
 para(h,'Heading 2');para(t)
refline(['mtr','park','mobility','weather','emergency'])
newpage();title('行前待确认清单','CHECK BEFORE DEPARTURE')
for i,issue in enumerate(D['issues'],1):para(f'{i}. {issue}')
para('建议领队优先处理的衔接','Heading 2')
para('先确认14日是否分组，以解除半小时重叠；再确认13日爱诗科技地址，复核30分钟转场。15日请接待方协助明确17W至15W的园内路径、登记方式，以及16:15结束访问后的上车安排。')
para('11日落地至晚宴仅95分钟，且包含机场手续；17日酒店出发至起飞175分钟，需覆盖赴机场交通和值机安检。两段都应结合当天航班与车况复核，不宜把表上的间隔直接当作交通时间。')
para('已确认与仍待确认的边界','Heading 2');para('“V4已列明”仅表示正式行程写明该活动和时刻；本路书未代替主办方再次向接待单位确认预约。开放网页提供的地址可用于核对楼宇，但不能代替领队的报到与入口通知。')
# Sources: requested anew supersedes old removal
newpage();title('资料来源与核验说明','SOURCES')
para('活动时间以香港校友会2026年10月7日V4正式行程为唯一依据；旧版手册仅作对照。本页列明信息核验来源，网址可直接点击。',size=9)
for s in D['sources']:
 p=para('',style='Caption')
 if s['url']:link(p,s['title'],s['url'])
 else:p.add_run(s['title'])
 p.paragraph_format.space_after=Pt(0)
para('地图坐标来自OpenStreetMap对象，使用Nominatim查询，并以官方地址核对；公爵大厦参考Commons大堂实景定位。完整坐标证据与照片署名见配套项目data.json及sources.json。地政总署接口本次未能取得可用响应，未将其写作核验成功的依据。',size=9)
# image credits 2 pages
photo_records=[]
seen=set()
for p in D['places']:
 if p['photo'] and p['image'] not in seen:seen.add(p['image']);photo_records.append((p['name'],p['photo']))
h=D['heroPhoto'];photo_records.insert(0,('封面 香港城市景观',dict(author=h['author'],source=h['source_page_url'],license=h['license'],licenseUrl=h['license_url'],caption=h['description'])))
for batch in range(math.ceil(len(photo_records)/7)):
 newpage();title('图片版权与署名'+(' 续' if batch else ''),'IMAGE CREDITS',toc_entry=batch==0)
 para('所有图片为真实地点的历史照片。适配中进行了尺寸缩放与局部显示裁切；照片继续适用原许可，下载及再利用请遵守来源页要求。',size=9)
 for name,a in photo_records[batch*7:batch*7+7]:
  para(name,'Heading 3');para(a['author']+' / Wikimedia Commons / '+a['license'],size=9)
  p=para('',style='Caption');link(p,'原始图片与文件信息',a['source']);p.add_run('  ·  ');link(p,'许可条款',a['licenseUrl'])
 para('地图底图与数据 © OpenStreetMap contributors。所有静态地图均为真实地理底图上的参考点，不绘制道路路线；在线地图底图使用OpenStreetMap服务。',size=9)
# Seed a real TOC field with cached contents. Page numbers updated from verified PDF if needed.
start=OxmlElement('w:r');fc=OxmlElement('w:fldChar');fc.set(qn('w:fldCharType'),'begin');start.append(fc);toc_holder._p.append(start)
r=OxmlElement('w:r');instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\f C \\h \\z ';r.append(instr);toc_holder._p.append(r)
r=OxmlElement('w:r');f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'separate');r.append(f);toc_holder._p.append(r)
prev=toc_holder._p
for text,pg,name in toc:
 p=OxmlElement('w:p');pr=OxmlElement('w:pPr');sp=OxmlElement('w:spacing');sp.set(qn('w:after'),'60');sp.set(qn('w:line'),'360');sp.set(qn('w:lineRule'),'auto');pr.append(sp);tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'right');tab.set(qn('w:leader'),'dot');tab.set(qn('w:pos'),'9500');tabs.append(tab);pr.append(tabs);p.append(pr)
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),name);r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=text;r.append(tx);h.append(r);p.append(h);r=OxmlElement('w:r');tab=OxmlElement('w:tab');r.append(tab);p.append(r);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' PAGEREF '+name+' \\h ');r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=str(pg);r.append(tx);f.append(r);p.append(f);prev.addnext(p);prev=p
r=OxmlElement('w:r');fc=OxmlElement('w:fldChar');fc.set(qn('w:fldCharType'),'end');r.append(fc);prev.append(r)
(R/'downloads').mkdir(exist_ok=True);out=R/'downloads/香港访学精美路书.docx';
for el in DOC.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
for el in DOC.styles.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
DOC.save(out);print('Saved',out,'planned pages',page)
(R/'.doc-build/page_plan.json').write_text(json.dumps(toc,ensure_ascii=False,indent=2))
